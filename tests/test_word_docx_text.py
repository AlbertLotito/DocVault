""".docx text lives in more places than body paragraphs.

python-docx's doc.paragraphs only covers top-level body paragraphs, so 31
library documents laid out in tables or text boxes (a price list, a school
physical form, a POST-code reference) failed with "No text found in document".
"""
import threading
import zipfile
from unittest.mock import MagicMock

from docx import Document

from extractors import microsoft_word_extractor as wx


def _ctx():
    from core.extractors.base import ExtractorContext, ExtractorLogger
    return ExtractorContext(vault_id='t', file_hash='t', cancel_token=threading.Event(),
                            logger=MagicMock(spec=ExtractorLogger), settings=MagicMock())


def test_body_paragraphs(tmp_path):
    doc = Document(); doc.add_paragraph('Hello from Word'); doc.add_paragraph('Second paragraph')
    p = tmp_path / 'a.docx'; doc.save(p)
    text, error = wx.extract(str(p), _ctx())
    assert error is None
    assert text.splitlines()[:3] == ['Hello from Word', '', 'Second paragraph'] or 'Second paragraph' in text


def test_table_only_document(tmp_path):
    doc = Document()
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = 'Código', 'Producto'
    t.cell(1, 0).text, t.cell(1, 1).text = 'A-17', 'Quinoa real'
    p = tmp_path / 'precios.docx'; doc.save(p)
    text, error = wx.extract(str(p), _ctx())
    assert error is None
    assert 'Código' in text and 'Quinoa real' in text


def test_headers_and_footers(tmp_path):
    doc = Document(); doc.add_paragraph('Body text')
    doc.sections[0].header.paragraphs[0].text = 'POCONO MOUNTAIN SCHOOL DISTRICT'
    doc.sections[0].footer.paragraphs[0].text = 'Form PH-2'
    p = tmp_path / 'form.docx'; doc.save(p)
    text, _ = wx.extract(str(p), _ctx())
    assert 'POCONO MOUNTAIN SCHOOL DISTRICT' in text and 'Form PH-2' in text and 'Body text' in text


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" ' \
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" ' \
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" ' \
    'xmlns:v="urn:schemas-microsoft-com:vml"'


def _raw_docx(path, body_xml):
    doc_xml = f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W}><w:body>{body_xml}</w:body></w:document>'
    with zipfile.ZipFile(path, 'w') as z:
        z.writestr('[Content_Types].xml',
                   '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                   '<Default Extension="xml" ContentType="application/xml"/>'
                   '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        z.writestr('word/document.xml', doc_xml)


def test_text_box_text_appears_once(tmp_path):
    """Text boxes are stored twice (mc:Choice + mc:Fallback); index them once."""
    box = '<w:txbxContent><w:p><w:r><w:t>LISTA DE PRECIOS SARIRI</w:t></w:r></w:p></w:txbxContent>'
    body = ('<w:p><w:r><mc:AlternateContent>'
            f'<mc:Choice Requires="wps"><wps:txbx>{box}</wps:txbx></mc:Choice>'
            f'<mc:Fallback><v:textbox>{box}</v:textbox></mc:Fallback>'
            '</mc:AlternateContent></w:r></w:p>')
    p = tmp_path / 'box.docx'; _raw_docx(p, body)
    text, error = wx.extract(str(p), _ctx())
    assert error is None
    assert text.count('LISTA DE PRECIOS SARIRI') == 1


def test_tabs_and_line_breaks(tmp_path):
    body = '<w:p><w:r><w:t>Code</w:t><w:tab/><w:t>Meaning</w:t><w:br/><w:t>1 short beep</w:t></w:r></w:p>'
    p = tmp_path / 'post.docx'; _raw_docx(p, body)
    text, _ = wx.extract(str(p), _ctx())
    assert 'Code\tMeaning\n1 short beep' in text


def test_really_empty_document_is_still_an_error(tmp_path):
    doc = Document(); p = tmp_path / 'empty.docx'; doc.save(p)
    text, error = wx.extract(str(p), _ctx())
    assert text is None and 'No text found' in error


def test_not_a_docx_is_an_error(tmp_path):
    p = tmp_path / 'fake.docx'; p.write_bytes(b'\x10Albert notes, plain bytes')
    text, error = wx.extract(str(p), _ctx())
    assert text is None and error.startswith('Failed to extract Word document')
