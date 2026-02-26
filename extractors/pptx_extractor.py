import os
from pptx import Presentation


def extract(file_path: str) -> tuple:
    print(f"  [pptx] Extracting: {os.path.basename(file_path)}")
    try:
        prs = Presentation(file_path)
        slides = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    t = shape.text_frame.text.strip()
                    if t:
                        texts.append(t)
            if texts:
                slides.append(f"[Slide {i}]\n" + "\n".join(texts))
        if not slides:
            return None, "No text found in presentation"
        return "\n\n".join(slides), None
    except Exception as e:
        return None, f"Failed to extract PowerPoint file: {e}"
