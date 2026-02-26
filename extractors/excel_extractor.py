import os
from openpyxl import load_workbook


def extract(file_path: str) -> tuple:
    print(f"  [excel] Extracting: {os.path.basename(file_path)}")
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        sections = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows():
                cells = [str(c.value) for c in row if c.value is not None]
                if cells:
                    rows.append("\t".join(cells))
            if rows:
                sections.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows))
        if not sections:
            return None, "No content found in spreadsheet"
        return "\n\n".join(sections), None
    except Exception as e:
        return None, f"Failed to extract Excel file: {e}"
